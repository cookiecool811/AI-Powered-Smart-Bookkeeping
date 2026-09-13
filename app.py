import os
import json
import re
import sqlite3
import base64
import requests
from io import BytesIO
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from datetime import datetime
from pathlib import Path
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from zhipuai import ZhipuAI
from dotenv import load_dotenv

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DATA_DIR / "expenses.db"
EXPORT_DIR = BASE_DIR / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)

CATEGORIES = [
    "餐饮", "交通", "购物", "住房", "娱乐", "医疗",
    "教育", "通讯", "旅行", "生活缴费", "数码电子", "其他"
]
INCOME_CATEGORIES = [
    "工资", "奖金", "投资收益", "兼职", "红包", "退款", "其他收入"
]

SYSTEM_PROMPT = f"""
你是一个专业的中文AI记账助手。
你的任务是把用户自然语言中的收支记录解析成结构化账单。
只返回合法JSON，不要Markdown，不要解释。

支出分类必须从以下类别中选择：
{", ".join(CATEGORIES)}

收入分类必须从以下类别中选择：
{", ".join(INCOME_CATEGORIES)}

JSON格式：
{{
  "items": [
    {{
      "amount": 12.5,
      "type": "expense",
      "category": "餐饮",
      "description": "午餐",
      "date": "YYYY-MM-DD",
      "payment_method": "未知",
      "confidence": 0.95
    }}
  ],
  "reply": "简短说明识别结果"
}}

规则：
1. amount 必须是正数。
2. type 为 "expense"（支出）或 "income"（收入），根据语义判断：工资、奖金、收入、收到、退款、红包等为收入；消费、买、花、付等为支出。
3. 如果用户没有明确日期，使用今天。
4. 一句话可能包含多笔记录，必须全部提取。
5. 如果没有金额，不要猜测，items 返回空数组。
6. 日期可以理解"今天、昨天、前天、上周"等表达。
7. 分类要根据语义判断，收入用收入分类，支出用支出分类。
8. confidence 为0到1之间的小数。
"""

def db():
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = db()
    cur = conn.cursor()
    cur.execute("""
    CREATE TABLE IF NOT EXISTS expenses (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        amount REAL NOT NULL,
        category TEXT NOT NULL,
        description TEXT NOT NULL,
        expense_date TEXT NOT NULL,
        payment_method TEXT DEFAULT '未知',
        confidence REAL DEFAULT 0,
        created_at TEXT NOT NULL
    )
    """)
    conn.commit()
    # Add type column if not exists (for existing databases)
    cur.execute("PRAGMA table_info(expenses)")
    cols = [row[1] for row in cur.fetchall()]
    if "type" not in cols:
        cur.execute("ALTER TABLE expenses ADD COLUMN type TEXT DEFAULT 'expense'")
        conn.commit()
    cur.close()
    conn.close()

def get_client():
    key = os.getenv("ZHIPUAI_API_KEY")
    if not key:
        raise RuntimeError("未配置 ZHIPUAI_API_KEY 环境变量")
    return ZhipuAI(api_key=key)

def parse_ai_json(text):
    text = text.strip()
    text = re.sub(r"^```json\s*", "", text)
    text = re.sub(r"^```\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return json.loads(text)

def ai_parse(user_text):
    client = get_client()
    today = datetime.now().strftime("%Y-%m-%d")
    prompt = f"今天是 {today}。请解析下面的消费记录：\n{user_text}"
    response = client.chat.completions.create(
        model=os.getenv("ZHIPUAI_MODEL", "glm-4.7-flash"),
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt}
        ],
        temperature=0.1,
        max_tokens=2000
    )
    content = response.choices[0].message.content
    data = parse_ai_json(content)
    if not isinstance(data, dict) or not isinstance(data.get("items"), list):
        raise ValueError("AI返回格式异常")
    return data

def save_items(items):
    conn = db()
    cur = conn.cursor()
    now = datetime.now().isoformat(timespec="seconds")
    saved = []
    for item in items:
        amount = float(item["amount"])
        if amount <= 0:
            continue
        item_type = item.get("type", "expense")
        if item_type not in ("expense", "income"):
            item_type = "expense"
        category = item.get("category", "其他")
        valid_cats = INCOME_CATEGORIES if item_type == "income" else CATEGORIES
        if category not in valid_cats:
            category = "其他收入" if item_type == "income" else "其他"
        description = str(item.get("description", "未命名消费"))
        date = str(item.get("date") or datetime.now().strftime("%Y-%m-%d"))
        payment = str(item.get("payment_method", "未知"))
        confidence = max(0, min(1, float(item.get("confidence", 0))))
        cur.execute("""
            INSERT INTO expenses
            (amount, type, category, description, expense_date, payment_method, confidence, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (amount, item_type, category, description, date, payment, confidence, now))
        new_id = cur.lastrowid
        saved.append({
            "id": new_id,
            "amount": amount,
            "type": item_type,
            "category": category,
            "description": description,
            "date": date,
            "payment_method": payment,
            "confidence": confidence
        })
    conn.commit()
    cur.close()
    conn.close()
    return saved

def query_expenses(start_date=None, end_date=None, ids=None):
    conn = db()
    cur = conn.cursor()
    sql = """
        SELECT id, amount, type, category, description,
               expense_date AS date, payment_method, confidence
        FROM expenses WHERE 1=1
    """
    params = []
    if ids:
        placeholders = ",".join("?" * len(ids))
        sql += " AND id IN (" + placeholders + ")"
        params.extend(ids)
    if start_date:
        sql += " AND expense_date >= ?"
        params.append(start_date)
    if end_date:
        sql += " AND expense_date <= ?"
        params.append(end_date)
    sql += " ORDER BY expense_date DESC, id DESC"
    cur.execute(sql, params)
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return [dict(r) for r in rows]

def resolve_date_range(range_type, start=None, end=None):
    today = datetime.now().strftime("%Y-%m-%d")
    if range_type == "month":
        return today[:7] + "-01", today
    if range_type == "year":
        return today[:4] + "-01-01", today
    if range_type == "custom":
        return (start or None), (end or None)
    return None, None

def range_label(range_type, start_date=None, end_date=None):
    if range_type == "month":
        return datetime.now().strftime("%Y年%m月")
    if range_type == "year":
        return datetime.now().strftime("%Y年")
    if range_type == "custom" and start_date and end_date:
        return f"{start_date}至{end_date}"
    return "全部"

def create_workbook(start_date=None, end_date=None, ids=None):
    rows = query_expenses(start_date, end_date, ids)
    wb = Workbook()
    ws = wb.active
    ws.title = "消费明细"

    headers = ["日期", "分类", "描述", "金额", "支付方式"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r["date"], r["category"], r["description"],
            r["amount"], r["payment_method"]
        ])

    total = round(sum(r["amount"] for r in rows), 2)
    ws.append(["合计", "", "", total, ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF2FF")

    for col, width in {"A":15, "B":14, "C":28, "D":14, "E":16}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    monthly = {}
    yearly = {}
    category = {}

    for r in rows:
        d = r["date"]
        month = d[:7]
        year = d[:4]
        monthly[month] = monthly.get(month, 0) + r["amount"]
        yearly[year] = yearly.get(year, 0) + r["amount"]
        category[r["category"]] = category.get(r["category"], 0) + r["amount"]

    def summary_sheet(name, mapping, first, second):
        sh = wb.create_sheet(name)
        sh.append([first, second])
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0F766E")
        for k in sorted(mapping.keys(), reverse=True):
            sh.append([k, round(mapping[k], 2)])
        sh.column_dimensions["A"].width = 20
        sh.column_dimensions["B"].width = 18

    summary_sheet("月度汇总", monthly, "月份", "消费总额")
    summary_sheet("年度汇总", yearly, "年份", "消费总额")
    summary_sheet("分类汇总", category, "消费分类", "消费总额")
    return wb

def build_excel():
    wb = create_workbook()
    filename = EXPORT_DIR / "AI记账消费报表.xlsx"
    wb.save(filename)
    return filename

def send_excel_email(start_date=None, end_date=None, ids=None):
    api_key = os.getenv("RESEND_API_KEY")
    mail_to = os.getenv("MAIL_TO")
    if not api_key or not mail_to:
        raise RuntimeError("未配置 RESEND_API_KEY 或 MAIL_TO 环境变量")
    wb = create_workbook(start_date, end_date, ids)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    excel_b64 = base64.b64encode(buf.read()).decode("utf-8")
    resp = requests.post(
        "https://api.resend.com/emails",
        headers={"Authorization": "Bearer " + api_key, "Content-Type": "application/json"},
        json={
            "from": "onboarding@resend.dev",
            "to": [mail_to],
            "subject": "AI记账消费报表 - " + datetime.now().strftime("%Y-%m-%d"),
            "text": "您好，附件是您的 AI 记账消费报表，请查收。",
            "attachments": [{"filename": "AI记账消费报表.xlsx", "content": excel_b64}]
        }
    )
    if resp.status_code != 200:
        raise RuntimeError("邮件发送失败: " + str(resp.status_code) + " " + resp.text)

@app.route("/")
def index():
    return render_template("index.html")

@app.get("/api/expenses")
def api_expenses():
    return jsonify({"items": query_expenses()})

@app.get("/api/stats")
def api_stats():
    rows = query_expenses()
    expenses = [r for r in rows if r.get("type", "expense") == "expense"]
    incomes = [r for r in rows if r.get("type") == "income"]
    month = datetime.now().strftime("%Y-%m")
    year = datetime.now().strftime("%Y")
    month_expense = sum(r["amount"] for r in expenses if r["date"].startswith(month))
    month_income = sum(r["amount"] for r in incomes if r["date"].startswith(month))
    year_expense = sum(r["amount"] for r in expenses if r["date"].startswith(year))
    year_income = sum(r["amount"] for r in incomes if r["date"].startswith(year))
    exp_cat = {}
    for r in expenses:
        exp_cat[r["category"]] = exp_cat.get(r["category"], 0) + r["amount"]
    exp_top = sorted(exp_cat.items(), key=lambda x: x[1], reverse=True)[:8]
    inc_cat = {}
    for r in incomes:
        inc_cat[r["category"]] = inc_cat.get(r["category"], 0) + r["amount"]
    inc_top = sorted(inc_cat.items(), key=lambda x: x[1], reverse=True)[:8]
    return jsonify({
        "total_expense": round(sum(r["amount"] for r in expenses), 2),
        "total_income": round(sum(r["amount"] for r in incomes), 2),
        "month_expense": round(month_expense, 2),
        "month_income": round(month_income, 2),
        "year_expense": round(year_expense, 2),
        "year_income": round(year_income, 2),
        "count": len(rows),
        "category": [{"name": k, "value": round(v, 2)} for k, v in exp_top],
        "income_category": [{"name": k, "value": round(v, 2)} for k, v in inc_top]
    })

@app.get("/api/summary")
def api_summary():
    range_type = request.args.get("range", "all")
    start = request.args.get("start")
    end = request.args.get("end")
    start_date, end_date = resolve_date_range(range_type, start, end)
    rows = query_expenses(start_date, end_date)
    total = round(sum(r["amount"] for r in rows), 2)
    return jsonify({
        "total": total,
        "count": len(rows),
        "range_label": range_label(range_type, start_date, end_date)
    })

@app.post("/api/parse")
def api_parse():
    data = request.get_json(silent=True) or {}
    text = str(data.get("text", "")).strip()
    if not text:
        return jsonify({"error": "请输入消费描述"}), 400
    try:
        result = ai_parse(text)
        items = save_items(result["items"])
        return jsonify({"reply": result.get("reply", "已完成记账"), "items": items})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.delete("/api/expenses/<int:expense_id>")
def api_delete(expense_id):
    conn = db()
    cur = conn.cursor()
    cur.execute("DELETE FROM expenses WHERE id=?", (expense_id,))
    conn.commit()
    cur.close()
    conn.close()
    return jsonify({"ok": True})

@app.get("/download/excel")
def download_excel():
    range_type = request.args.get("range", "current")
    start = request.args.get("start")
    end = request.args.get("end")
    ids = request.args.getlist("ids")
    if ids:
        ids = [int(x) for x in ids]
    if range_type == "current":
        start_date, end_date = None, None
    else:
        start_date, end_date = resolve_date_range(range_type, start, end)
    wb = create_workbook(start_date, end_date, ids)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"AI记账消费报表_{range_label(range_type, start_date, end_date)}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.post("/api/send-email")
def api_send_email():
    data = request.get_json(silent=True) or {}
    range_type = data.get("range", "current")
    start = data.get("start")
    end = data.get("end")
    ids = data.get("ids")
    if range_type == "current":
        start_date, end_date = None, None
    else:
        start_date, end_date = resolve_date_range(range_type, start, end)
    try:
        send_excel_email(start_date, end_date, ids)
        return jsonify({"ok": True, "message": "报表已发送到邮箱"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.get("/health")
def health():
    return jsonify({"status": "ok"})

init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=True)
